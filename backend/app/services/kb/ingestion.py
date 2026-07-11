import io
import logging
import uuid
from datetime import datetime, timezone
from pathlib import Path

from sqlalchemy.ext.asyncio import AsyncSession

from app.db.models import KBChunk, KBDocument
from app.services.llm.base import BaseLLMProvider

logger = logging.getLogger(__name__)

import re

CHUNK_SIZE = 500  # tokens approx
CHUNK_OVERLAP = 50
CJK_MAX_CHARS = 260  # chars per chunk for CJK/space-poor text
CJK_OVERLAP = 40

_SENT_SPLIT = re.compile(r"(?<=[。！？!?；;\n])")


def _split_text(text: str, chunk_size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
    """
    Chunk text for embedding. Latin/space-rich text is split on word
    boundaries; CJK or space-poor text is split on sentence punctuation and
    packed to ~CJK_MAX_CHARS, since ``str.split()`` would collapse an entire
    Chinese paragraph into a single unusable chunk.
    """
    words = text.split()
    # Heuristic: if words are on average long, the text has few spaces (CJK).
    space_poor = len(words) < len(text) / 6

    if not space_poor:
        chunks = []
        i = 0
        while i < len(words):
            chunks.append(" ".join(words[i: i + chunk_size]))
            i += chunk_size - overlap
        return [c for c in chunks if c.strip()]

    # CJK path: sentence-aware packing
    pieces = [p.strip() for p in _SENT_SPLIT.split(text) if p.strip()]
    chunks: list[str] = []
    cur = ""
    for p in pieces:
        if len(cur) + len(p) <= CJK_MAX_CHARS:
            cur += p
        else:
            if cur:
                chunks.append(cur)
            while len(p) > CJK_MAX_CHARS:
                chunks.append(p[:CJK_MAX_CHARS])
                p = p[CJK_MAX_CHARS - CJK_OVERLAP:]
            cur = p
    if cur:
        chunks.append(cur)
    return [c for c in chunks if c.strip()]


def _extract_text(content: bytes, filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".txt":
        return content.decode("utf-8", errors="replace")
    elif suffix == ".pdf":
        try:
            import PyPDF2
            reader = PyPDF2.PdfReader(io.BytesIO(content))
            return "\n".join(p.extract_text() or "" for p in reader.pages)
        except Exception as e:
            logger.warning("PDF parse failed: %s", e)
            return ""
    elif suffix in (".docx",):
        try:
            import docx
            doc = docx.Document(io.BytesIO(content))
            return "\n".join(p.text for p in doc.paragraphs)
        except Exception as e:
            logger.warning("DOCX parse failed: %s", e)
            return ""
    elif suffix in (".xlsx", ".xls"):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
            lines = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    lines.append(" | ".join(str(c) for c in row if c is not None))
            return "\n".join(lines)
        except Exception as e:
            logger.warning("XLSX parse failed: %s", e)
            return ""
    return content.decode("utf-8", errors="replace")


async def ingest_document(
    db: AsyncSession,
    document: KBDocument,
    content: bytes,
    embed_provider: BaseLLMProvider,
) -> int:
    """Extract text, chunk, embed, and store. Returns number of chunks created."""
    text = _extract_text(content, document.title)
    if not text.strip():
        logger.warning("No text extracted from document %s", document.id)
        return 0

    chunks_text = _split_text(text)
    chunk_objs = []

    # Batch embed
    BATCH = 50
    all_embeddings: list[list[float]] = []
    for i in range(0, len(chunks_text), BATCH):
        batch = chunks_text[i: i + BATCH]
        try:
            resp = await embed_provider.embed(batch)
            all_embeddings.extend(resp.embeddings)
        except Exception as e:
            logger.error("Embed batch failed: %s", e)
            all_embeddings.extend([None] * len(batch))

    for idx, (chunk_text, embedding) in enumerate(zip(chunks_text, all_embeddings)):
        chunk = KBChunk(
            id=uuid.uuid4(),
            tenant_id=document.tenant_id,
            document_id=document.id,
            group_id=document.group_id,
            content=chunk_text,
            chunk_index=idx,
            embedding=embedding,
            token_count=len(chunk_text.split()),
        )
        db.add(chunk)

    document.chunk_count = len(chunks_text)
    document.status = "ready"
    await db.commit()
    return len(chunks_text)
